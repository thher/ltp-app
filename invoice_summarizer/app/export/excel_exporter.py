"""Excel export — 6-sheet workbook using openpyxl."""
from __future__ import annotations

from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.export.base_exporter import BaseExporter

# ── Palette ────────────────────────────────────────────────────────────────────
_HDR_BG = "0F172A"
_HDR_FG = "FFFFFF"
_ALT_BG = "F0F4F8"
_WHITE  = "FFFFFF"
_ACCENT = "1E40AF"
_BORDER = "CBD5E1"
_LABEL  = "334155"


class ExcelExporter(BaseExporter):
    """Produces a .xlsx workbook with six sheets."""

    def export(self, data: dict, output_path: Path) -> Path:
        out = Path(output_path)
        if out.suffix != ".xlsx":
            out = out.with_suffix(".xlsx")

        wb = openpyxl.Workbook()
        wb.remove(wb.active)  # drop the blank default sheet

        self._sheet_summary(wb, data)
        self._sheet_invoices(wb, data.get("invoices", []))
        self._sheet_suppliers(wb, data.get("suppliers", []))
        self._sheet_products(wb, data.get("products", []))
        self._sheet_material_categories(wb, data.get("material_categories", []))
        self._sheet_review_queue(wb, data.get("review_queue", []))

        wb.save(out)
        return out

    # ── sheets ────────────────────────────────────────────────────────────────

    def _sheet_summary(self, wb, data: dict) -> None:
        ws = wb.create_sheet("Summary")
        ws.column_dimensions["A"].width = 28
        ws.column_dimensions["B"].width = 30

        kpis    = data.get("kpis", {})
        filters = data.get("filters", {})
        ts      = data.get("generated_at", datetime.now())

        rows = [
            ("Invoice Summarizer — Export Report", None, "title"),
            ("Generated", str(ts)[:19],            "meta"),
            (None,        None,                    "blank"),
            ("Filters Applied", None,              "section"),
            ("Date From",     filters.get("date_from") or "All", "meta"),
            ("Date To",       filters.get("date_to")   or "All", "meta"),
            ("Supplier ID",   filters.get("supplier_id") or "All", "meta"),
            ("Material Category", filters.get("material_category") or "All", "meta"),
            (None, None, "blank"),
            ("Key Performance Indicators", None, "section"),
            ("Active Suppliers",  kpis.get("active_suppliers", "—"), "meta"),
            ("Total Invoices",    kpis.get("invoice_count",    "—"), "meta"),
            ("Total Spend",       _fmt(kpis.get("total_spend")),     "meta"),
            ("Average Invoice",   _fmt(kpis.get("avg_invoice")),     "meta"),
        ]

        for r_idx, (label, value, style) in enumerate(rows, 1):
            if style == "blank":
                continue
            ca = ws.cell(r_idx, 1, label)
            cb = ws.cell(r_idx, 2, value) if value is not None else None

            if style == "title":
                ca.font = Font(bold=True, size=16, color=_HDR_FG)
                ca.fill = PatternFill("solid", fgColor=_HDR_BG)
                ca.alignment = Alignment(horizontal="left")
                ws.row_dimensions[r_idx].height = 28
                ws.merge_cells(f"A{r_idx}:B{r_idx}")
            elif style == "section":
                ca.font = Font(bold=True, size=11, color=_HDR_FG)
                ca.fill = PatternFill("solid", fgColor=_ACCENT)
                ws.merge_cells(f"A{r_idx}:B{r_idx}")
            elif style == "meta":
                ca.font = Font(bold=True, color=_LABEL)
                if cb:
                    cb.alignment = Alignment(horizontal="left")

    def _sheet_invoices(self, wb, rows: list[dict]) -> None:
        _write_table(
            wb.create_sheet("Invoices"),
            headers   = ["Invoice #", "Date", "Due Date", "Supplier",
                         "Currency", "Subtotal", "VAT", "Grand Total", "Status"],
            keys      = ["invoice_number", "invoice_date", "due_date", "supplier_name",
                         "currency", "subtotal", "vat_total", "grand_total", "status"],
            rows      = rows,
            num_cols  = {5, 6, 7},
            col_widths= [16, 12, 12, 28, 10, 16, 12, 16, 12],
        )

    def _sheet_suppliers(self, wb, rows: list[dict]) -> None:
        _write_table(
            wb.create_sheet("Suppliers"),
            headers   = ["Supplier", "Category", "Country", "Invoices", "Total Spend"],
            keys      = ["canonical_name", "category_name", "country",
                         "invoice_count", "total_spend"],
            rows      = rows,
            num_cols  = {4},
            col_widths= [32, 22, 14, 10, 16],
        )

    def _sheet_products(self, wb, rows: list[dict]) -> None:
        _write_table(
            wb.create_sheet("Products"),
            headers   = ["Supplier", "Description", "Category", "Unit",
                         "Total Qty", "Length (m)", "Total Spend", "Count", "Review?"],
            keys      = ["supplier_name", "raw_description", "material_category", "unit",
                         "total_quantity", "total_length_m", "total_spend",
                         "occurrences", "needs_review"],
            rows      = rows,
            num_cols  = {4, 5, 6},
            col_widths= [24, 38, 16, 8, 12, 12, 16, 8, 10],
        )

    def _sheet_material_categories(self, wb, rows: list[dict]) -> None:
        _write_table(
            wb.create_sheet("Material Categories"),
            headers   = ["Category", "Unit Type", "Total Quantity", "Total Spend", "Items"],
            keys      = ["material_category", "unit_type",
                         "total_quantity", "total_spend", "item_count"],
            rows      = rows,
            num_cols  = {2, 3},
            col_widths= [22, 14, 16, 16, 10],
        )

    def _sheet_review_queue(self, wb, rows: list[dict]) -> None:
        _write_table(
            wb.create_sheet("Review Queue"),
            headers   = ["ID", "Supplier", "Invoice #", "Issue Type",
                         "Description", "Suggestion", "Resolved"],
            keys      = ["id", "supplier_name", "invoice_number", "issue_type",
                         "description", "suggestion", "resolved"],
            rows      = rows,
            col_widths= [6, 24, 16, 20, 36, 36, 10],
        )


# ── helpers ───────────────────────────────────────────────────────────────────

def _write_table(
    ws,
    headers: list[str],
    keys: list[str],
    rows: list[dict],
    num_cols: set[int] | None = None,
    col_widths: list[int] | None = None,
) -> None:
    num_cols = num_cols or set()
    thin = Side(style="thin", color=_BORDER)
    brd  = Border(left=thin, right=thin, top=thin, bottom=thin)

    # Header row
    for c, hdr in enumerate(headers, 1):
        cell = ws.cell(1, c, hdr)
        cell.font      = Font(bold=True, color=_HDR_FG, size=9)
        cell.fill      = PatternFill("solid", fgColor=_HDR_BG)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border    = brd
    ws.row_dimensions[1].height = 20
    ws.freeze_panes = "A2"

    # Data rows
    for r, row in enumerate(rows, 2):
        bg = PatternFill("solid", fgColor=(_ALT_BG if r % 2 == 0 else _WHITE))
        for c, key in enumerate(keys, 1):
            raw = row.get(key)
            if raw is None:
                val: object = ""
            elif isinstance(raw, bool):
                val = "Yes" if raw else "No"
            else:
                val = raw
            cell = ws.cell(r, c, val)
            cell.fill   = bg
            cell.border = brd
            if (c - 1) in num_cols and isinstance(val, (int, float)):
                cell.number_format = "#,##0.00"
                cell.alignment = Alignment(horizontal="right")
            else:
                cell.alignment = Alignment(horizontal="left", wrap_text=False)

    # Column widths
    if col_widths:
        for i, w in enumerate(col_widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w


def _fmt(v) -> str:
    if v is None:
        return "—"
    try:
        return f"{float(v):,.2f}"
    except (TypeError, ValueError):
        return str(v)
