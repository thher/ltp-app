"""Tests for ExportService, ExcelExporter, and PDFReportExporter."""
from __future__ import annotations

import sys
from datetime import date
from pathlib import Path

import pytest

ROOT = Path(__file__).parent.parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from app.database.db_manager import DatabaseManager
from app.database.migrations import MigrationRunner
from app.database.models import Invoice, LineItem, Supplier
from app.database.repositories.aggregation_repo import AggregationRepository
from app.database.repositories.invoice_repo import InvoiceRepository
from app.database.repositories.line_item_repo import LineItemRepository
from app.database.repositories.supplier_repo import SupplierRepository
from app.export.excel_exporter import ExcelExporter
from app.export.export_service import ExportService
from app.export.pdf_exporter import PDFReportExporter


# ── fixtures ──────────────────────────────────────────────────────────────────

@pytest.fixture
def db():
    mgr = DatabaseManager(db_path=Path(":memory:"))
    MigrationRunner(mgr).run()
    yield mgr
    mgr.close()


@pytest.fixture
def export_service(db):
    return ExportService(db, AggregationRepository(db))


@pytest.fixture
def populated_db(db):
    """DB with one supplier, two invoices, and four line items."""
    sup_repo  = SupplierRepository(db)
    inv_repo  = InvoiceRepository(db)
    li_repo   = LineItemRepository(db)

    sup = sup_repo.save(Supplier(
        canonical_name="Nydal Bygg AS",
        normalized_key="nydal bygg as",
        country="NO",
    ))
    inv1 = inv_repo.save(Invoice(
        supplier_id=sup.id,
        invoice_number="INV-001",
        invoice_date=date(2025, 3, 15),
        currency="NOK",
        grand_total=12_500.00,
        status="processed",
        original_path="/tmp/inv1.pdf",
        copy_path="/tmp/inv1_copy.pdf",
        file_hash="abc123",
    ))
    inv2 = inv_repo.save(Invoice(
        supplier_id=sup.id,
        invoice_number="INV-002",
        invoice_date=date(2025, 6, 20),
        currency="NOK",
        grand_total=8_300.00,
        status="processed",
        original_path="/tmp/inv2.pdf",
        copy_path="/tmp/inv2_copy.pdf",
        file_hash="def456",
    ))
    for inv, desc, cat, unit in [
        (inv1, "48X148 IMPREGNERT LEKTER",   "Timber",  "m"),
        (inv1, "TRANSPORT KRANBIL",           "Transport","stk"),
        (inv2, "28X120 ROYAL TERRASSEBORD",  "Decking", "m"),
        (inv2, "GLAVA A-PLATE 10CM",          "Insulation","m²"),
    ]:
        li_repo.save(LineItem(
            invoice_id=inv.id,
            raw_description=desc,
            quantity=10.0,
            unit=unit,
            line_total=2_500.00,
            unit_type=unit,
            normalized_quantity=10.0,
            material_category=cat,
            confidence=1.0,
        ))
    return db


@pytest.fixture
def populated_service(populated_db):
    return ExportService(populated_db, AggregationRepository(populated_db))


# ── ExportService tests ───────────────────────────────────────────────────────

class TestExportService:
    def test_returns_all_keys(self, export_service):
        data = export_service.gather_data()
        expected = {"generated_at", "filters", "kpis", "invoices",
                    "suppliers", "products", "material_categories", "review_queue"}
        assert expected.issubset(data.keys())

    def test_empty_db_returns_empty_lists(self, export_service):
        data = export_service.gather_data()
        assert data["invoices"] == []
        assert data["suppliers"] == []
        assert data["products"] == []
        assert data["material_categories"] == []
        assert data["review_queue"] == []

    def test_filters_recorded(self, export_service):
        data = export_service.gather_data(
            date_from=date(2025, 1, 1),
            date_to=date(2025, 12, 31),
            material_category="Timber",
        )
        assert data["filters"]["date_from"] == "2025-01-01"
        assert data["filters"]["date_to"]   == "2025-12-31"
        assert data["filters"]["material_category"] == "Timber"

    def test_populated_data(self, populated_service):
        data = populated_service.gather_data()
        assert len(data["invoices"])  == 2
        assert len(data["suppliers"]) == 1
        assert len(data["products"])  == 4

    def test_date_filter_restricts_invoices(self, populated_service):
        data = populated_service.gather_data(
            date_from=date(2025, 6, 1),
            date_to=date(2025, 6, 30),
        )
        assert len(data["invoices"]) == 1
        assert data["invoices"][0]["invoice_number"] == "INV-002"

    def test_material_category_filter(self, populated_service):
        data = populated_service.gather_data(material_category="Timber")
        # Only the Timber line item should appear in products
        assert all(
            r["material_category"] == "Timber"
            for r in data["products"]
        )

    def test_kpis_totals(self, populated_service):
        data = populated_service.gather_data()
        kpis = data["kpis"]
        assert kpis["invoice_count"] == 2
        assert abs(float(kpis["total_spend"]) - 20_800.0) < 0.01


# ── ExcelExporter tests ───────────────────────────────────────────────────────

class TestExcelExporter:
    def test_creates_xlsx(self, export_service, tmp_path):
        data = export_service.gather_data()
        out  = ExcelExporter().export(data, tmp_path / "report")
        assert out.exists()
        assert out.suffix == ".xlsx"

    def test_adds_xlsx_suffix(self, export_service, tmp_path):
        data = export_service.gather_data()
        out  = ExcelExporter().export(data, tmp_path / "report.csv")
        assert out.suffix == ".xlsx"

    def test_has_six_sheets(self, export_service, tmp_path):
        import openpyxl
        data = export_service.gather_data()
        out  = ExcelExporter().export(data, tmp_path / "report.xlsx")
        wb   = openpyxl.load_workbook(out)
        assert set(wb.sheetnames) == {
            "Summary", "Invoices", "Suppliers",
            "Products", "Material Categories", "Review Queue",
        }

    def test_invoices_sheet_has_data(self, populated_service, tmp_path):
        import openpyxl
        data = populated_service.gather_data()
        out  = ExcelExporter().export(data, tmp_path / "report.xlsx")
        wb   = openpyxl.load_workbook(out)
        ws   = wb["Invoices"]
        # Header + 2 data rows
        assert ws.max_row == 3

    def test_products_sheet_has_data(self, populated_service, tmp_path):
        import openpyxl
        data = populated_service.gather_data()
        out  = ExcelExporter().export(data, tmp_path / "report.xlsx")
        wb   = openpyxl.load_workbook(out)
        ws   = wb["Products"]
        assert ws.max_row == 5   # header + 4 line items

    def test_summary_title_in_cell_a1(self, export_service, tmp_path):
        import openpyxl
        data = export_service.gather_data()
        out  = ExcelExporter().export(data, tmp_path / "report.xlsx")
        wb   = openpyxl.load_workbook(out)
        ws   = wb["Summary"]
        assert "Invoice Summarizer" in (ws["A1"].value or "")

    def test_material_categories_sheet(self, populated_service, tmp_path):
        import openpyxl
        data = populated_service.gather_data()
        out  = ExcelExporter().export(data, tmp_path / "report.xlsx")
        wb   = openpyxl.load_workbook(out)
        ws   = wb["Material Categories"]
        # header row + at least one category row
        assert ws.max_row >= 2


# ── PDFReportExporter tests ───────────────────────────────────────────────────

class TestPDFReportExporter:
    def test_creates_pdf(self, export_service, tmp_path):
        data = export_service.gather_data()
        out  = PDFReportExporter().export(data, tmp_path / "report")
        assert out.exists()
        assert out.suffix == ".pdf"

    def test_adds_pdf_suffix(self, export_service, tmp_path):
        data = export_service.gather_data()
        out  = PDFReportExporter().export(data, tmp_path / "report.xlsx")
        assert out.suffix == ".pdf"

    def test_non_empty_pdf(self, export_service, tmp_path):
        data = export_service.gather_data()
        out  = PDFReportExporter().export(data, tmp_path / "report.pdf")
        assert out.stat().st_size > 1_000

    def test_populated_pdf_larger(self, populated_service, tmp_path):
        data = populated_service.gather_data()
        out  = PDFReportExporter().export(data, tmp_path / "report.pdf")
        assert out.stat().st_size > 5_000

    def test_pdf_is_valid_pdf(self, export_service, tmp_path):
        data = export_service.gather_data()
        out  = PDFReportExporter().export(data, tmp_path / "report.pdf")
        with open(out, "rb") as f:
            header = f.read(5)
        assert header == b"%PDF-"
